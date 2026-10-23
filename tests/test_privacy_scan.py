import tempfile
import unittest
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import privacy_scan as ps

class ScanTests(unittest.TestCase):
    def test_underscore_placeholder(self):
        got=ps.scan_blanks([ps.Line("合同价款：人民币________元", "2", "文本层", "第2页")])
        self.assertEqual(len(got),1)
        self.assertEqual(got[0].category,"下划线/占位符")
        self.assertEqual(got[0].risk,"高")

    def test_empty_label_after_colon(self):
        got=ps.scan_blanks([ps.Line("项目名称：", "1", "文本层", "第1页")])
        self.assertEqual(len(got),1)
        self.assertEqual(got[0].category,"疑似空白字段标签")

    def test_table_border_noise_ignored(self):
        got=ps.scan_blanks([ps.Line("+----------------+", "1", "文本层", "第1页")])
        self.assertEqual(got,[])

    def test_privacy_redaction(self):
        got=ps.scan_privacy([ps.Line("联系电话：13812345678", "3", "文本层", "第3页")])
        self.assertEqual(len(got),1)
        self.assertNotIn("13812345678", got[0].evidence)
        self.assertIn("****", got[0].value)

    def test_checkbox_candidate(self):
        got=ps.scan_blanks([ps.Line("选项 □ 适用", "4", "文本层", "第4页")])
        self.assertEqual(len(got),1)
        self.assertEqual(got[0].category,"未选方框（待核实）")

    def test_docx_formatted_underline_whitespace(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"blank.docx"
            doc=Document()
            p=doc.add_paragraph("合同编号：")
            run=p.add_run("     ")
            run.underline=True
            doc.save(path)
            lines=ps.read_docx(path)
            got=ps.scan_blanks(lines)
            self.assertTrue(any(f.category=="DOCX格式下划线空白" for f in got))

if __name__ == "__main__":
    unittest.main()
