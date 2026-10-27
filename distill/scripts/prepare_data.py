"""Extract source documents into auditable DeepSeek teacher tasks."""
from __future__ import annotations
import argparse, hashlib, re
from pathlib import Path
from .common import DATA, ROOT, write_jsonl

SUPPORTED = {".pdf", ".docx", ".txt", ".html", ".htm"}

def extract(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        import fitz
        doc = fitz.open(path)
        return [{"page": i + 1, "text": page.get_text("text").strip()} for i, page in enumerate(doc) if page.get_text("text").strip()]
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    chunks.append(" | ".join(cells))
        return [{"page": None, "text": "\n".join(chunks)}] if chunks else []
    text = path.read_text(encoding="utf-8", errors="replace")
    if ext in {".html", ".htm"}:
        from html.parser import HTMLParser
        class Strip(HTMLParser):
            def __init__(self):
                super().__init__(); self.parts=[]
            def handle_data(self, data):
                if data.strip(): self.parts.append(data.strip())
        parser=Strip(); parser.feed(text); text="\n".join(parser.parts)
    return [{"page": None, "text": text.strip()}] if text.strip() else []

def main() -> int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--input-dir", default=str(ROOT / "distill" / "data" / "input"))
    ap.add_argument("--output", default=str(DATA / "tasks.jsonl"))
    ap.add_argument("--max-chars", type=int, default=12000, help="每份材料每个片段最大字符数")
    args=ap.parse_args()
    input_dir=Path(args.input_dir)
    if not input_dir.exists():
        print(f"输入目录不存在：{input_dir}。请建立 distill/data/input/<项目ID>/ 并放入招标文件、合同、公示。")
        return 2
    tasks=[]
    for project in sorted(p for p in input_dir.iterdir() if p.is_dir()):
        documents=[]
        for path in sorted(p for p in project.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED):
            try:
                pages=extract(path)
            except Exception as exc:
                print(f"[失败] {path}: {type(exc).__name__}: {exc}")
                continue
            for page in pages:
                text=page["text"]
                # 保留来源；超长内容切块。跨文件对比需在生成阶段明确标记上下文不完整。
                for offset in range(0, len(text), args.max_chars):
                    chunk=text[offset:offset+args.max_chars]
                    if not chunk.strip(): continue
                    documents.append({"source":str(path.relative_to(input_dir)), "page":page["page"], "chunk_index":offset//args.max_chars, "text":chunk})
        if not documents:
            print(f"[跳过] {project.name}: 没有提取到文本")
            continue
        digest=hashlib.sha256("\n".join(d["text"] for d in documents).encode()).hexdigest()
        tasks.append({"task_id":f"{project.name}-{digest[:12]}","project_id":project.name,"task_type":"contract_cross_check","documents":documents,
          "instruction":"请分别提取合同关键字段，并比较招标文件、合同、中标公示。只引用输入中可见的证据；公示未披露的内容标记为未披露，不推断一致。"})
    write_jsonl(Path(args.output), tasks)
    print(f"已生成 {len(tasks)} 个项目任务：{args.output}")
    return 0 if tasks else 1

if __name__ == "__main__":
    raise SystemExit(main())
