#!/usr/bin/env python3
"""合同隐私与空白项初筛工具。启发式结果必须人工复核。"""
from __future__ import annotations
import argparse, re, sys
from dataclasses import dataclass
from pathlib import Path

try:
    from docx import Document
except ImportError:
    Document = None

@dataclass
class Line:
    text: str
    page: str = "未标页"
    source: str = "文本层"
    location: str = ""

@dataclass
class Finding:
    category: str
    value: str
    page: str
    location: str
    source: str
    risk: str
    evidence: str
    note: str = ""

UNDERSCORE_RE = re.compile(r"[_＿﹍﹎﹏]{3,}")
EMPTY_BRACKET_RE = re.compile(r"[（(]\s*[）)]|\[\s*\]|【\s*】")
UNCHECKED_RE = re.compile(r"□|☐|▢")
FILL_PROMPT_RE = re.compile(r"填写|此处注明|待定|待补充|另行填写|由双方填写|以实际填写为准|请填写|待确认")
LABEL_EMPTY_RE = re.compile(r"(?:项目名称|合同编号|签订地点|联系人|联系电话|项目负责人|工期|服务期限|质保期|质量标准|履约保证金|质保金比例|违约金率|收款账户|开户行|送达地址|附件清单|合同价款|签订日期)\s*[:：]\s*$")
PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d(?:[\s-]?\d){8}(?!\d)")
ID_RE = re.compile(r"(?<![0-9Xx])\d{17}[0-9Xx](?![0-9Xx])|(?<!\d)\d{15}(?!\d)")
QQ_CONTEXT_RE = re.compile(r"(?:QQ号|QQ号码|QQ\s*[:：])\s*\D{0,8}(\d{5,12})", re.I)
BANK_CONTEXT_RE = re.compile(r"(?:银行账号|银行卡号|收款账号|收款账户|开户账号|个人账户|账号)\s*[:：]?\s*([\d -]{8,32})")
SIGNING_TERMS = re.compile(r"签字|签名|盖章|公章|法定代表人或授权代表|落款日期|签署日期|签订地点")
HIGH_TERMS = re.compile(r"金额|价款|合同价|工期|服务期限|起算|质保金|违约金|担保|履约保证金|争议解决|管辖|收款账户|项目负责人|质量标准")
MEDIUM_TERMS = re.compile(r"联系人|联系电话|送达地址|通知地址|交付地点|附件清单|开户行")

def mask_digits(raw: str, kind: str) -> str:
    digits = re.sub(r"\D", "", raw)
    if kind == "phone": return digits[:3] + "****" + digits[-4:] if len(digits) >= 7 else "[已脱敏]"
    if kind == "id": return digits[:3] + "*" * max(0, len(digits)-7) + digits[-4:] if len(digits) >= 15 else "[已脱敏]"
    return digits[:4] + "****" + digits[-4:] if len(digits) >= 8 else "[已脱敏]"

def redact(text: str) -> str:
    text = PHONE_RE.sub(lambda m: mask_digits(m.group(), "phone"), text)
    text = ID_RE.sub(lambda m: mask_digits(m.group(), "id"), text)
    text = re.sub(r"(?i)(QQ号|QQ号码|QQ\s*[:：]\s*)(\d{5,12})", lambda m: m.group(1)+mask_digits(m.group(2), "account"), text)
    text = BANK_CONTEXT_RE.sub(lambda m: m.group(0)[:len(m.group(0))-len(m.group(1))]+mask_digits(m.group(1), "account"), text)
    return text

def read_pdf(path: Path, layout: bool=False):
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise SystemExit("缺少 PDF 文本依赖，请运行：pip install pypdf") from e
    lines=[]; reader=PdfReader(str(path))
    for pageno,page in enumerate(reader.pages,1):
        for line in (page.extract_text() or "").splitlines():
            lines.append(Line(line,str(pageno),"文本层",f"第{pageno}页"))
    visual=[]
    if layout:
        try: import pdfplumber
        except ImportError:
            print("提示：未安装 pdfplumber，跳过 PDF 版式线检测。可运行 pip install pdfplumber",file=sys.stderr)
        else:
            try:
                with pdfplumber.open(str(path)) as pdf:
                    for pageno,page in enumerate(pdf.pages,1):
                        chars=page.chars or []
                        for obj in (page.lines or [])+(page.rects or []):
                            x0=float(obj.get("x0",0)); x1=float(obj.get("x1",0)); top=float(obj.get("top",0)); bottom=float(obj.get("bottom",top))
                            if x1-x0<28 or abs(bottom-top)>3.5: continue
                            nearby=[c for c in chars if abs(float(c.get("top",0))-top)<=7 and float(c.get("x1",0))>=x0-2 and float(c.get("x0",0))<=x1+2]
                            if nearby: continue
                            visual.append(Finding("疑似版式填空线","[需查看页面图像]",str(pageno),f"第{pageno}页，横线约 x={x0:.0f}–{x1:.0f}","版式层","待复核","检测到较长且附近无文字的横线；也可能是表格边框或装饰线。","需结合页面图像判断是否为未填写字段。"))
            except Exception as e: print(f"提示：PDF 版式检测失败，已跳过：{e}",file=sys.stderr)
    return lines,visual

def iter_docx_cells(table):
    seen=set()
    for ri,row in enumerate(table.rows,1):
        for ci,cell in enumerate(row.cells,1):
            key=id(cell._tc)
            if key in seen: continue
            seen.add(key); yield ri,ci,cell

def read_docx(path: Path):
    if Document is None: raise SystemExit("缺少 DOCX 依赖，请运行：pip install python-docx")
    doc=Document(str(path)); lines=[]
    for idx,p in enumerate(doc.paragraphs,1):
        loc=f"正文段落 {idx}"; lines.append(Line(p.text,"未标页","文本层",loc))
        for ri,run in enumerate(p.runs,1):
            if run.underline and run.text and not run.text.strip(): lines.append(Line(p.text + " [WORD_UNDERLINE_WHITESPACE]","未标页","DOCX格式层",f"{loc}，run {ri} 下划线空白"))
    for ti,table in enumerate(doc.tables,1):
        for ri,ci,cell in iter_docx_cells(table):
            loc=f"表格 {ti} 第{ri}行第{ci}列"; lines.append(Line(cell.text,"未标页","文本层",loc))
            if not cell.text.strip(): lines.append(Line("[EMPTY_TABLE_CELL]","未标页","DOCX表格层",loc))
            for pi,p in enumerate(cell.paragraphs,1):
                for rj,run in enumerate(p.runs,1):
                    if run.underline and run.text and not run.text.strip(): lines.append(Line(p.text + " [WORD_UNDERLINE_WHITESPACE]","未标页","DOCX格式层",f"{loc}，段落{pi} run {rj} 下划线空白"))
    xml=doc._element.xml
    if "w:ffData" in xml or "w:fldChar" in xml: lines.append(Line("[DOCX_FORM_FIELD_XML]","未标页","DOCX XML层","文档含表单域/域代码标记，需复核是否已填写"))
    return lines

def read_input(path: Path, layout=False):
    ext=path.suffix.lower()
    if ext==".pdf": return read_pdf(path,layout)
    if ext==".docx": return read_docx(path),[]
    if ext in (".txt",".md"): return [Line(s,"未标页","文本层",f"第{i}行") for i,s in enumerate(path.read_text(encoding="utf-8",errors="replace").splitlines(),1)],[]
    if ext==".doc": raise SystemExit("不支持直接读取旧版 .doc；请先另存为 .docx 或 PDF。")
    raise SystemExit(f"不支持的格式：{ext}")

def risk_for(context,category):
    if SIGNING_TERMS.search(context): return "低（可能为正常签署留白）"
    if HIGH_TERMS.search(context): return "高"
    if MEDIUM_TERMS.search(context): return "中"
    if category in ("空白括号/方框","未填写提示语","疑似空白字段标签","空白表格单元格","DOCX格式下划线空白","DOCX表单域"): return "待复核（按字段用途定级）"
    return "低/待复核"

def scan_blanks(lines):
    out=[]; seen=set()
    for line in lines:
        text=line.text; cat=None
        if "[WORD_UNDERLINE_WHITESPACE]" in text: cat="DOCX格式下划线空白"
        elif text=="[EMPTY_TABLE_CELL]": cat="空白表格单元格"
        elif text=="[DOCX_FORM_FIELD_XML]": cat="DOCX表单域"
        elif UNDERSCORE_RE.search(text):
            if re.search(r"(?:\+[-＿_]{3,}|\|\s*[-＿_]{3,}|[-＿_]{5,}\s*\|)",text): continue
            cat="下划线/占位符"
        elif EMPTY_BRACKET_RE.search(text): cat="空白括号/方框"
        elif UNCHECKED_RE.search(text): cat="未选方框（待核实）"
        elif FILL_PROMPT_RE.search(text): cat="未填写提示语"
        elif LABEL_EMPTY_RE.search(text): cat="疑似空白字段标签"
        if not cat: continue
        key=(line.page,line.location,cat,text)
        if key in seen: continue
        seen.add(key); note=""
        if cat=="空白表格单元格": note="空单元格可能表示无/不适用，不自动认定为缺陷。"
        elif cat=="DOCX格式下划线空白": note="Word 下划线格式应用于空白文本，需核对字段语义。"
        elif cat=="未选方框（待核实）": note="需确认是否应勾选，不能仅凭未选判错。"
        out.append(Finding(cat,"[不适用]",line.page,line.location,line.source,risk_for(text+" "+line.location,cat),redact(text),note))
    return out

def scan_privacy(lines):
    out=[]; seen=set()
    for line in lines:
        matches=[]
        matches += [("手机号",mask_digits(m.group(),"phone")) for m in PHONE_RE.finditer(line.text)]
        matches += [("身份证号",mask_digits(m.group(),"id")) for m in ID_RE.finditer(line.text)]
        matches += [("QQ号",mask_digits(m.group(1),"account")) for m in QQ_CONTEXT_RE.finditer(line.text)]
        for m in BANK_CONTEXT_RE.finditer(line.text):
            raw=m.group(1).strip()
            if len(re.sub(r"\D","",raw))>=8: matches.append(("疑似银行账号",mask_digits(raw,"account")))
        for cat,masked in matches:
            key=(line.page,line.location,cat,masked)
            if key in seen: continue
            seen.add(key); out.append(Finding(cat,masked,line.page,line.location,line.source,"需人工核实",redact(line.text),"已脱敏显示；请核实是否为个人信息及是否确有必要出现在合同中。"))
    return out

def tender_blank_compare(tender_path: Path, contract_findings):
    """保守字段级提示，不代替条款语义对齐。"""
    tender_lines,_=read_input(tender_path,False); tender_text="\n".join(x.text for x in tender_lines)
    labels=["项目名称","合同编号","签订地点","联系人","联系电话","项目负责人","工期","服务期限","质保期","质量标准","履约保证金","质保金比例","违约金率","收款账户","开户行","送达地址","附件清单","合同价款","签订日期"]
    out=[]
    for label in labels:
        if label in tender_text:
            for f in contract_findings:
                if label in f.evidence or label in f.location:
                    out.append(Finding("招标字段待核对",f.value,f.page,f.location,f.source,f.risk,f.evidence,f"招标文件出现字段“{label}”，合同对应处疑似留空；需人工确认条款对应关系。"))
    return out

def write_report(path,mode,privacy,blanks,layout,tender):
    allb=blanks+layout+tender
    risk_counts={r:sum(f.risk==r for f in allb) for r in sorted({f.risk for f in allb})}
    cat_counts={c:sum(f.category==c for f in allb) for c in sorted({f.category for f in allb})}
    out=[f"# 合同初筛报告：{path.name}","","自动初筛仅供复核，不是法律意见或完整合规结论。扫描版 PDF 需 OCR；旧版 DOC 需转换。","","## 统计摘要",f"- 扫描模式：{mode}",f"- 隐私候选：{len(privacy)}",f"- 空白候选：{len(allb)}","- 按风险："+("；".join(f"{k}：{v}" for k,v in risk_counts.items()) or "无"),"- 按类别："+("；".join(f"{k}：{v}" for k,v in cat_counts.items()) or "无"),"","## 隐私信息候选（已脱敏）","| 序号 | 类别 | 脱敏值 | 页码/位置 | 来源 | 说明 |","|---:|---|---|---|---|---|"]
    if privacy:
        for i,f in enumerate(privacy,1): out.append(f"| {i} | {f.category} | {f.value} | {f.page} / {f.location} | {f.source} | {f.note} |")
    else: out.append("| — | 未发现匹配项 | — | — | — | 不代表绝对不存在隐私信息。 |")
    out += ["","## 合同未填写及缺失信息候选","| 序号 | 类别 | 页码/位置 | 原文片段（已脱敏） | 风险等级 | 来源层 | 说明/建议 |","|---:|---|---|---|---|---|---|"]
    grouped=[]; seen=set()
    for f in sorted(allb,key=lambda x:(0 if x.risk=="高" else 1 if x.risk=="中" else 2,x.page,x.location,x.category)):
        key=(f.page,f.location,f.category,f.evidence)
        if key in seen: continue
        seen.add(key); grouped.append(f)
    if grouped:
        for i,f in enumerate(grouped,1):
            ev=redact(f.evidence).replace("|","\\|")
            out.append(f"| {i} | {f.category} | {f.page} / {f.location} | {ev} | {f.risk} | {f.source} | {f.note or '结合上下文判断是否必填；未知内容标记待合同双方确认。'} |")
    else: out.append("| — | 未发现明显候选项 | — | — | — | — | 仍需复核版式和 OCR 限制。 |")
    out += ["","## 限制与人工复核","- 下划线、空白单元格、未选方框和表单域可能是正常模板格式，不能自动等同于缺陷。","- 签字、盖章、签署日期的空白可能属于正常签署留白。","- PDF 版式线检测会把表格边框/装饰线作为候选，需查看页面图像。","- 招标字段匹配是保守提示，不等同于条款级语义对齐。","- 所有引用上下文必须脱敏；OCR、图片、手写和格式变形可能漏检。"]
    return "\n".join(out)+"\n"

def main():
    ap=argparse.ArgumentParser(description="合同隐私信息与空白项初筛")
    ap.add_argument("file",help="合同文件 PDF/DOCX/TXT/MD")
    ap.add_argument("--out",help="Markdown 报告路径")
    ap.add_argument("--mode",choices=["all","privacy","blanks"],default="all")
    ap.add_argument("--layout",action="store_true",help="使用可选 pdfplumber 做 PDF 横线初筛")
    ap.add_argument("--tender",help="可选招标文件路径，提供字段级交叉提示")
    args=ap.parse_args(); path=Path(args.file); lines,layout=read_input(path,args.layout)
    privacy=scan_privacy(lines) if args.mode in ("all","privacy") else []
    blanks=scan_blanks(lines) if args.mode in ("all","blanks") else []
    tender=tender_blank_compare(Path(args.tender),blanks) if args.tender and blanks else []
    report=write_report(path,args.mode,privacy,blanks,layout if args.mode in ("all","blanks") else [],tender)
    if args.out: Path(args.out).write_text(report,encoding="utf-8"); print(f"报告已写入：{args.out}")
    else: print(report)
if __name__=="__main__": main()
